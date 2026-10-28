# -*- coding: utf-8 -*-
"""T321 Excel Dashboard - independent development module; does not replace the stable T321 app."""
import os, sys, math, statistics, traceback
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import openpyxl
    from openpyxl.chart import BarChart, Reference, LineChart, PieChart
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None
try:
    import matplotlib
    matplotlib.use('TkAgg')
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
except ImportError:
    Figure = None
try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

APP_TITLE = 'دستیار اکسل T321 | داشبورد هوشمند'


def as_number(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else None
    if isinstance(v, str):
        s = v.strip().replace(',', '').replace('٬', '').replace('٫', '.')
        trans = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
        s = s.translate(trans)
        try:
            n = float(s)
            return n if math.isfinite(n) else None
        except Exception:
            return None
    return None


def display_num(n):
    if n is None: return '—'
    if abs(n - round(n)) < 1e-9: return f'{int(round(n)):,}'
    return f'{n:,.2f}'


def shape_fa(s):
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(str(s)))
    except Exception:
        return str(s)


class DashboardApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry('1180x800')
        self.root.minsize(920, 620)
        self.path = ''
        self.wb = None
        self.headers = []
        self.rows = []
        self.results = []
        self.chart_canvas = None
        self._style()
        self._build()

    def _style(self):
        st = ttk.Style()
        try: st.theme_use('clam')
        except Exception: pass
        st.configure('TButton', font=('Tahoma', 10), padding=7)
        st.configure('TLabel', font=('Tahoma', 10))
        st.configure('Title.TLabel', font=('Tahoma', 17, 'bold'), foreground='#263a66')
        st.configure('KPI.TLabel', font=('Tahoma', 12, 'bold'), foreground='#243b64')
        st.configure('Treeview', font=('Tahoma', 9), rowheight=26)
        st.configure('Treeview.Heading', font=('Tahoma', 9, 'bold'))

    def _build(self):
        header = ttk.Frame(self.root, padding=(14, 12))
        header.pack(fill='x')
        ttk.Label(header, text=APP_TITLE, style='Title.TLabel').pack(side='right')
        ttk.Label(header, text='نسخه آزمایشی مستقل ـ فایل اصلی را تغییر نمی‌دهد', foreground='#68758a').pack(side='left')

        top = ttk.LabelFrame(self.root, text='۱) انتخاب فایل و معیار تحلیل', padding=10)
        top.pack(fill='x', padx=14, pady=(0, 8))
        ttk.Button(top, text='باز کردن فایل اکسل', command=self.open_file).grid(row=0, column=0, padx=5, pady=5, sticky='ew')
        self.file_label = ttk.Label(top, text='هنوز فایلی انتخاب نشده است', width=48, anchor='e')
        self.file_label.grid(row=0, column=1, padx=6, sticky='ew')
        ttk.Label(top, text='برگه:').grid(row=0, column=2, padx=4)
        self.sheet_var = tk.StringVar()
        self.sheet_box = ttk.Combobox(top, textvariable=self.sheet_var, state='readonly', width=18)
        self.sheet_box.grid(row=0, column=3, padx=4)
        self.sheet_box.bind('<<ComboboxSelected>>', lambda e: self.load_sheet())
        ttk.Label(top, text='ستون عنوان/گروه:').grid(row=1, column=0, padx=4, pady=5, sticky='e')
        self.category_var = tk.StringVar()
        self.category_box = ttk.Combobox(top, textvariable=self.category_var, state='readonly', width=25)
        self.category_box.grid(row=1, column=1, padx=5, sticky='ew')
        ttk.Label(top, text='ستون عددی:').grid(row=1, column=2, padx=4, sticky='e')
        self.value_var = tk.StringVar()
        self.value_box = ttk.Combobox(top, textvariable=self.value_var, state='readonly', width=18)
        self.value_box.grid(row=1, column=3, padx=4)
        ttk.Label(top, text='معیار بهترین:').grid(row=2, column=0, padx=4, pady=5, sticky='e')
        self.best_mode = tk.StringVar(value='بیشترین مقدار بهتر است')
        self.best_box = ttk.Combobox(top, textvariable=self.best_mode, state='readonly', width=25, values=['بیشترین مقدار بهتر است','کمترین مقدار بهتر است'])
        self.best_box.grid(row=2, column=1, padx=5, sticky='ew')
        ttk.Button(top, text='ساخت داشبورد', command=self.analyze).grid(row=2, column=3, padx=8)
        for c in (1,): top.columnconfigure(c, weight=1)

        kpis = ttk.Frame(self.root, padding=(14, 0, 14, 8))
        kpis.pack(fill='x')
        self.kpi_labels = []
        for title in ('تعداد ردیف معتبر', 'مجموع', 'میانگین', 'بیشترین مقدار', 'کمترین مقدار'):
            card = ttk.LabelFrame(kpis, text=title, padding=10)
            card.pack(side='right', fill='both', expand=True, padx=4)
            val = ttk.Label(card, text='—', style='KPI.TLabel', anchor='center')
            val.pack(fill='x')
            self.kpi_labels.append(val)

        body = ttk.Panedwindow(self.root, orient='horizontal')
        body.pack(fill='both', expand=True, padx=14, pady=(0, 8))
        left = ttk.LabelFrame(body, text='رتبه‌بندی و جزئیات', padding=6)
        right = ttk.LabelFrame(body, text='نمودار مقایسه', padding=6)
        body.add(left, weight=1); body.add(right, weight=1)
        self.tree = ttk.Treeview(left, columns=('rank','name','value'), show='headings')
        for col, label, width in [('rank','رتبه',55),('name','عنوان',220),('value','مقدار',110)]:
            self.tree.heading(col, text=label)
            self.tree.column(col, width=width, anchor='center' if col != 'name' else 'e')
        yscroll = ttk.Scrollbar(left, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=yscroll.set)
        self.tree.pack(side='right', fill='both', expand=True); yscroll.pack(side='left', fill='y')
        self.chart_frame = right
        self.chart_placeholder = ttk.Label(right, text='پس از انتخاب فایل و ستون‌ها، داشبورد اینجا نمایش داده می‌شود', anchor='center', justify='center')
        self.chart_placeholder.pack(fill='both', expand=True)

        bottom = ttk.Frame(self.root, padding=(14, 0, 14, 12))
        bottom.pack(fill='x')
        ttk.Button(bottom, text='خروجی اکسل داشبورددار', command=self.export_excel).pack(side='right', padx=4)
        ttk.Button(bottom, text='گزارش PDF قابل چاپ', command=self.export_pdf).pack(side='right', padx=4)
        ttk.Button(bottom, text='راهنما', command=self.show_help).pack(side='left', padx=4)
        self.status = ttk.Label(bottom, text='آماده', foreground='#52617a')
        self.status.pack(side='left', padx=12)

    def open_file(self):
        if openpyxl is None:
            messagebox.showerror('نیازمندی', 'کتابخانه openpyxl نصب نیست. از بسته کامل برنامه استفاده کنید.')
            return
        p = filedialog.askopenfilename(title='انتخاب فایل اکسل', filetypes=[('Excel files','*.xlsx *.xlsm'), ('All files','*.*')])
        if not p: return
        try:
            self.path = p
            self.wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
            self.sheet_box['values'] = self.wb.sheetnames
            self.sheet_var.set(self.wb.sheetnames[0])
            self.file_label.config(text=os.path.basename(p))
            self.load_sheet()
            self.status.config(text='فایل با موفقیت خوانده شد')
        except Exception as e:
            messagebox.showerror('خطا در خواندن فایل', str(e))

    def load_sheet(self):
        if not self.wb or not self.sheet_var.get(): return
        ws = self.wb[self.sheet_var.get()]
        data = ws.iter_rows(values_only=True)
        first = next(data, None)
        if not first:
            messagebox.showwarning('برگه خالی', 'این برگه داده‌ای ندارد.'); return
        headers = []
        for i, v in enumerate(first):
            name = str(v).strip() if v is not None and str(v).strip() else f'ستون {i+1}'
            headers.append(name)
        self.headers = headers
        self.rows = [list(r) for r in data if any(v is not None for v in r)]
        self.category_box['values'] = headers
        self.value_box['values'] = headers
        self.category_var.set(headers[0])
        numeric = None
        for idx, h in enumerate(headers):
            if any(as_number(row[idx] if idx < len(row) else None) is not None for row in self.rows[:50]):
                numeric = h; break
        self.value_var.set(numeric or headers[0])
        self.results = []
        self.clear_results()
        self.status.config(text=f'{len(self.rows)} ردیف بارگذاری شد')

    def clear_results(self):
        for item in self.tree.get_children(): self.tree.delete(item)
        for k in self.kpi_labels: k.config(text='—')
        if self.chart_canvas:
            self.chart_canvas.get_tk_widget().destroy(); self.chart_canvas = None
        self.chart_placeholder = ttk.Label(self.chart_frame, text='برای مشاهده نمودار، «ساخت داشبورد» را بزنید', anchor='center')
        self.chart_placeholder.pack(fill='both', expand=True)

    def analyze(self):
        if not self.rows or not self.headers:
            messagebox.showinfo('انتخاب فایل', 'ابتدا یک فایل اکسل باز کنید.'); return
        ci, vi = self.headers.index(self.category_var.get()), self.headers.index(self.value_var.get())
        groups = {}
        valid = 0
        for row in self.rows:
            raw = row[vi] if vi < len(row) else None
            num = as_number(raw)
            if num is None: continue
            name = row[ci] if ci < len(row) else None
            name = str(name).strip() if name is not None and str(name).strip() else '(بدون عنوان)'
            groups[name] = groups.get(name, 0.0) + num
            valid += 1
        if not groups:
            messagebox.showwarning('داده عددی پیدا نشد', 'ستون عددی انتخاب‌شده مقدار قابل تحلیل ندارد. ستون دیگری انتخاب کنید.'); return
        self.results = sorted(groups.items(), key=lambda x: x[1], reverse=(self.best_mode.get() == 'بیشترین مقدار بهتر است'))
        vals = [as_number(r[vi] if vi < len(r) else None) for r in self.rows]
        vals = [v for v in vals if v is not None]
        total = sum(vals); avg = total / len(vals) if vals else 0
        self.kpi_labels[0].config(text=f'{valid:,}')
        self.kpi_labels[1].config(text=display_num(total))
        self.kpi_labels[2].config(text=display_num(avg))
        self.kpi_labels[3].config(text=display_num(max(vals) if vals else None))
        self.kpi_labels[4].config(text=display_num(min(vals) if vals else None))
        for item in self.tree.get_children(): self.tree.delete(item)
        for rank, (name, val) in enumerate(self.results, 1):
            self.tree.insert('', 'end', values=(rank, name, display_num(val)))
        self.draw_chart()
        self.status.config(text=f'تحلیل انجام شد؛ {len(self.results)} گروه')

    def draw_chart(self):
        if Figure is None:
            messagebox.showwarning('نمودار', 'برای نمایش نمودار، matplotlib باید نصب باشد.'); return
        if self.chart_canvas: self.chart_canvas.get_tk_widget().destroy(); self.chart_canvas = None
        if self.chart_placeholder: self.chart_placeholder.destroy(); self.chart_placeholder = None
        top = self.results[:10]
        fig = Figure(figsize=(6.2, 4.4), dpi=100)
        ax = fig.add_subplot(111)
        names = [x[0] for x in reversed(top)]
        values = [x[1] for x in reversed(top)]
        ax.barh(range(len(names)), values)
        ax.set_yticks(range(len(names)))
        ax.set_yticklabels([n[:22] for n in names], fontsize=8)
        ax.set_title(('۱۰ مورد برتر — ' if self.best_mode.get() == 'بیشترین مقدار بهتر است' else '۱۰ مقدار کم‌تر — ') + self.value_var.get(), fontsize=10)
        ax.grid(axis='x', alpha=.2)
        fig.tight_layout()
        self.chart_canvas = FigureCanvasTkAgg(fig, master=self.chart_frame)
        self.chart_canvas.draw(); self.chart_canvas.get_tk_widget().pack(fill='both', expand=True)

    def export_excel(self):
        if not self.results:
            messagebox.showinfo('ابتدا تحلیل کنید', 'ابتدا فایل را باز و «ساخت داشبورد» را اجرا کنید.'); return
        p = filedialog.asksaveasfilename(title='ذخیره داشبورد اکسل', defaultextension='.xlsx', initialfile='T321_Dashboard.xlsx', filetypes=[('Excel Workbook','*.xlsx')])
        if not p: return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active; ws.title = 'داشبورد'
            ws.sheet_view.rightToLeft = True
            ws.merge_cells('A1:D1'); ws['A1'] = 'داشبورد هوشمند T321'
            ws['A1'].font = Font(name='Tahoma', size=18, bold=True, color='FFFFFF')
            ws['A1'].fill = PatternFill('solid', fgColor='263A66')
            ws['A1'].alignment = Alignment(horizontal='center')
            ws.append(['ستون گروه‌بندی', self.category_var.get(), 'ستون مقدار', self.value_var.get()])
            ws.append(['معیار رتبه‌بندی', self.best_mode.get()])
            nums = [as_number(r[self.headers.index(self.value_var.get())] if self.headers.index(self.value_var.get()) < len(r) else None) for r in self.rows]
            nums = [n for n in nums if n is not None]
            metrics = [('تعداد ردیف معتبر',len(nums)),('مجموع',sum(nums)),('میانگین',sum(nums)/len(nums) if nums else 0),('بیشترین مقدار',max(nums) if nums else 0),('کمترین مقدار',min(nums) if nums else 0)]
            ws.append(['شاخص','مقدار'])
            for row in metrics: ws.append(list(row))
            start = 10
            ws.cell(start,1,'رتبه'); ws.cell(start,2,'عنوان'); ws.cell(start,3,'مقدار')
            for i, (name, val) in enumerate(self.results, start+1):
                ws.cell(i,1,i-start); ws.cell(i,2,name); ws.cell(i,3,val)
            chart = BarChart(); chart.type = 'bar'; chart.style = 10; chart.title = 'رتبه‌بندی بر اساس ' + self.value_var.get(); chart.y_axis.title = self.category_var.get(); chart.x_axis.title = self.value_var.get(); chart.height = 9; chart.width = 18
            data = Reference(ws, min_col=3, min_row=start, max_row=start+min(len(self.results),20))
            cats = Reference(ws, min_col=2, min_row=start+1, max_row=start+min(len(self.results),20))
            chart.add_data(data, titles_from_data=True); chart.set_categories(cats); ws.add_chart(chart, 'F3')
            for col, width in {'A':25,'B':30,'C':18,'D':25}.items(): ws.column_dimensions[col].width = width
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(horizontal='right' if cell.column in (2,4) else 'center', vertical='center')
                    if cell.row in (2,3,start):
                        cell.font = Font(name='Tahoma', bold=True, color='FFFFFF')
                        cell.fill = PatternFill('solid', fgColor='45649A')
            raw = wb.create_sheet('داده‌های منبع')
            raw.sheet_view.rightToLeft = True
            raw.append(self.headers)
            for r in self.rows: raw.append(r)
            for col in range(1, len(self.headers)+1): raw.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
            wb.save(p)
            messagebox.showinfo('انجام شد', f'فایل داشبورد ذخیره شد:\n{p}')
        except Exception as e:
            messagebox.showerror('خطا در خروجی اکسل', f'{e}\n\n{traceback.format_exc()}')

    def export_pdf(self):
        if not self.results:
            messagebox.showinfo('ابتدا تحلیل کنید', 'ابتدا فایل را باز و «ساخت داشبورد» را اجرا کنید.'); return
        if not REPORTLAB_OK:
            messagebox.showerror('نیازمندی PDF', 'کتابخانه reportlab نصب نیست. از بسته کامل برنامه استفاده کنید.'); return
        p = filedialog.asksaveasfilename(title='ذخیره گزارش PDF', defaultextension='.pdf', initialfile='T321_Report.pdf', filetypes=[('PDF','*.pdf')])
        if not p: return
        try:
            doc = SimpleDocTemplate(p, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
            styles = getSampleStyleSheet()
            font_path = r'C:\Windows\Fonts\tahoma.ttf'
            if os.path.exists(font_path):
                pdfmetrics.registerFont(TTFont('TahomaT321', font_path)); font_name = 'TahomaT321'
            else:
                font_name = 'Helvetica'
            title_style = ParagraphStyle('T321Title', parent=styles['Title'], fontName=font_name, alignment=TA_CENTER, fontSize=16, leading=22)
            normal = ParagraphStyle('T321Normal', parent=styles['Normal'], fontName=font_name, alignment=TA_RIGHT, fontSize=9, leading=14)
            story = [Paragraph(shape_fa('گزارش داشبورد هوشمند T321'), title_style), Spacer(1, 10)]
            story.append(Paragraph(shape_fa(f'تاریخ گزارش: {datetime.now().strftime("%Y-%m-%d %H:%M")}'), normal))
            story.append(Paragraph(shape_fa(f'گروه‌بندی: {self.category_var.get()} | معیار: {self.value_var.get()} | {self.best_mode.get()}'), normal)); story.append(Spacer(1, 10))
            vi = self.headers.index(self.value_var.get())
            nums = [as_number(r[vi] if vi < len(r) else None) for r in self.rows]
            nums = [n for n in nums if n is not None]
            metrics = [['شاخص','مقدار'], ['تعداد معتبر',str(len(nums))], ['مجموع',display_num(sum(nums))], ['میانگین',display_num(sum(nums)/len(nums) if nums else 0)], ['بیشترین',display_num(max(nums) if nums else None)], ['کمترین',display_num(min(nums) if nums else None)]]
            t = Table([[shape_fa(str(c)) for c in row] for row in metrics], colWidths=[180,180])
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#263A66')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.grey),('ALIGN',(0,0),(-1,-1),'RIGHT'),('FONTNAME',(0,0),(-1,-1),font_name),('FONTSIZE',(0,0),(-1,-1),9),('PADDING',(0,0),(-1,-1),6)]))
            story += [t, Spacer(1, 12), Paragraph(shape_fa('رتبه‌بندی گروه‌ها'), normal), Spacer(1, 5)]
            rank_data = [['رتبه','عنوان','مقدار']]
            for i, (name, val) in enumerate(self.results[:40],1): rank_data.append([str(i), shape_fa(name), display_num(val)])
            rt = Table(rank_data, colWidths=[55,350,120], repeatRows=1)
            rt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#45649A')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.35,colors.grey),('ALIGN',(0,0),(-1,-1),'RIGHT'),('FONTNAME',(0,0),(-1,-1),font_name),('FONTSIZE',(0,0),(-1,-1),8),('PADDING',(0,0),(-1,-1),5)]))
            story.append(rt)
            doc.build(story)
            messagebox.showinfo('انجام شد', f'گزارش PDF ذخیره شد:\n{p}')
        except Exception as e:
            messagebox.showerror('خطا در PDF', f'{e}\n\nدر صورت مشکل نمایش فارسی، باید فونت و شکل‌دهی فارسی در نسخه ویندوز آزمایش شود.')

    def show_help(self):
        messagebox.showinfo('راهنمای T321', '۱) فایل xlsx یا xlsm را باز کنید.\n۲) برگه و ستون عنوان و عددی را انتخاب کنید.\n۳) ساخت داشبورد را بزنید.\n۴) خروجی اکسل یا PDF بگیرید.\n\nاین نسخه ماژول آزمایشی داشبورد است و هنوز جایگزین برنامه اصلی T321 نیست.')


def main():
    if openpyxl is None:
        root = tk.Tk(); root.withdraw(); messagebox.showerror('نیازمندی', 'openpyxl نصب نیست. لطفاً از نسخه بسته‌بندی‌شده استفاده کنید.'); return
    root = tk.Tk()
    DashboardApp(root)
    root.mainloop()

if __name__ == '__main__':
    main()
